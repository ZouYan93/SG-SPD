import os
import re

def count_model_parameters(model):
    total_params = sum(p.numel() for p in model.parameters())
    trainable_params = sum(p.numel() for p in model.parameters() if p.requires_grad)
    return total_params, trainable_params


def get_file_size_mb(file_path):
    if not os.path.exists(file_path):
        return None
    return os.path.getsize(file_path) / (1024 * 1024)


def _get_workbook(excel_path):
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise ImportError(
            'openpyxl 库未安装，请运行 "pip install openpyxl" 来安装该库，以使用 Excel 相关功能。'
        ) from exc

    parent_dir = os.path.dirname(excel_path)
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)

    if os.path.exists(excel_path):
        workbook = load_workbook(excel_path)
    else:
        workbook = Workbook()
    return workbook


def _sanitize_sheet_name(sheet_name):
    # Excel 不允许: : \\ / ? * [ ]，且长度最大 31
    sanitized = re.sub(r'[:\\/?*\[\]]', '_', str(sheet_name))
    sanitized = sanitized.strip() or 'sheet'
    return sanitized[:31]


def _append_kv_section(sheet, title, data):
    if title:
        sheet.append([title])
    for key, value in data.items():
        sheet.append([key, value])
    sheet.append([])


def _append_table_section(sheet, title, headers, rows):
    if title:
        sheet.append([title])
    if headers:
        sheet.append(list(headers))
    for row in rows:
        if isinstance(row, dict):
            sheet.append([row.get(h, None) for h in headers])
        else:
            sheet.append(list(row))
    sheet.append([])


def write_sections_to_excel(excel_path, sheet_name, sections, mode='append'):
    """
    sections: list[dict]
      - {'type': 'kv', 'title': str, 'data': dict}
      - {'type': 'table', 'title': str, 'headers': list[str], 'rows': list[dict|list]}
    mode:
      - 'append': 追加到现有 sheet 尾部
      - 'replace': 清空并重建该 sheet
    """
    workbook = _get_workbook(excel_path)
    sanitized_sheet_name = _sanitize_sheet_name(sheet_name)

    if sanitized_sheet_name in workbook.sheetnames:
        if mode == 'replace':
            old_sheet = workbook[sanitized_sheet_name]
            workbook.remove(old_sheet)
            sheet = workbook.create_sheet(sanitized_sheet_name)
        else:
            sheet = workbook[sanitized_sheet_name]
            if sheet.max_row > 0:
                sheet.append([])
    else:
        sheet = workbook.create_sheet(sanitized_sheet_name)

    # 若默认 sheet 为空且新建了业务 sheet，则移除默认 sheet，避免空白表
    if 'Sheet' in workbook.sheetnames and len(workbook.sheetnames) > 1:
        default_sheet = workbook['Sheet']
        if default_sheet.max_row == 1 and default_sheet.max_column == 1 and default_sheet['A1'].value is None:
            workbook.remove(default_sheet)

    for section in sections:
        section_type = section.get('type')
        if section_type == 'kv':
            _append_kv_section(sheet, section.get('title', ''), section.get('data', {}))
        elif section_type == 'table':
            _append_table_section(
                sheet,
                section.get('title', ''),
                section.get('headers', []),
                section.get('rows', []),
            )
        else:
            raise ValueError(f"不支持的 section type: {section_type}")

    workbook.save(excel_path)